"""Mock VAHAN vehicle registry.

data/mock_vahan_registry.xlsx (scripts/generate_registry.py: 100 fictitious
vehicles) plus private rows built from real test documents, which are kept
out of the public repo: data/private_registry_rows.json (gitignored) or the
REGISTRY_PRIVATE_ROWS env var (a JSON list). load() puts them all in the
vehicle_registry table at startup, on either backend.
"""
import json
import os
import re
from datetime import date, datetime
from pathlib import Path

import db

ROOT = Path(__file__).resolve().parent
REGISTRY_FILE = ROOT / "data" / "mock_vahan_registry.xlsx"
PRIVATE_FILE = ROOT / "data" / "private_registry_rows.json"
STANDARD_PLATE = re.compile(r"^[A-Z]{2}\d{1,2}[A-Z]{0,3}\d{4}$")
BH_PLATE = re.compile(r"^\d{2}BH\d{4}[A-Z]{1,2}$")
INT_COLUMNS = {"idv", "ncb_percent"}


def today():
    """The registry's notion of today; tests pin it."""
    return date.today()


def normalise_plate(text):
    return re.sub(r"[\s\-.]", "", str(text or "")).upper()


def valid_plate(plate):
    plate = normalise_plate(plate)
    return bool(STANDARD_PLATE.match(plate) or BH_PLATE.match(plate))


def _cell(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return value


def _spreadsheet_rows():
    if not REGISTRY_FILE.exists():
        print(f"[registry] {REGISTRY_FILE.name} not found; registry is empty", flush=True)
        return []
    from openpyxl import load_workbook

    sheet = load_workbook(REGISTRY_FILE, read_only=True).active
    rows = sheet.iter_rows(values_only=True)
    header = [str(name) for name in next(rows)]
    return [dict(zip(header, (_cell(value) for value in row))) for row in rows if row and row[0]]


def _private_rows():
    raw = os.getenv("REGISTRY_PRIVATE_ROWS", "").strip()
    if not raw and PRIVATE_FILE.exists():
        raw = PRIVATE_FILE.read_text(encoding="utf-8")
    if not raw:
        return []
    try:
        rows = json.loads(raw)
    except ValueError:
        print("[registry] private rows are not valid JSON; ignored", flush=True)
        return []
    return [row for row in rows if isinstance(row, dict)]


def records():
    by_plate = {}
    for row in _spreadsheet_rows() + _private_rows():
        record = {key: row.get(key) for key in row}
        record["plate"] = normalise_plate(record.get("plate"))
        for column in INT_COLUMNS:
            try:
                record[column] = int(record.get(column) or 0)
            except (TypeError, ValueError):
                record[column] = 0
        record["owner_mobile"] = str(record.get("owner_mobile") or "")
        if record["plate"]:
            by_plate[record["plate"]] = record
    return list(by_plate.values())


def load(force=False):
    """Fill the vehicle_registry table. Skipped when it already holds the
    same number of rows, so repeated startups cost one count query."""
    rows = records()
    if not force and rows and db.registry_count() == len(rows):
        return len(rows)
    db.registry_replace(rows)
    print(f"[registry] loaded {len(rows)} vehicles", flush=True)
    return len(rows)


def lookup(plate):
    """The registration record for a plate, or None.

    Production: this is where the VAHAN (Parivahan) API call goes, by
    registration number, returning the same fields. Here it reads the mock
    registry loaded by load().
    """
    plate = normalise_plate(plate)
    if not valid_plate(plate):
        return None
    try:
        record = db.registry_get(plate)
        if record is None and db.registry_count() == 0:
            load()
            record = db.registry_get(plate)
    except Exception as error:
        # An unavailable registry is "not checked", never a failed claim.
        print(f"[registry] lookup unavailable: {type(error).__name__}", flush=True)
        return None
    return record


def _day(value):
    try:
        return date.fromisoformat(str(value)[:10])
    except (TypeError, ValueError):
        return None


def policy_state(record, on=None):
    """active | expired | not_started | third_party_only | unknown, on a date."""
    if not record:
        return "unknown"
    if record.get("policy_type") == "third_party_only":
        return "third_party_only"
    on = on or today()
    start, end = _day(record.get("policy_start")), _day(record.get("policy_end"))
    if not start or not end:
        return "unknown"
    if on < start:
        return "not_started"
    if on > end:
        return "expired"
    return "active"


def policy_covers(record, day):
    """True/False when the dates are known; None when they are not."""
    start, end = _day((record or {}).get("policy_start")), _day((record or {}).get("policy_end"))
    if not start or not end or not day:
        return None
    return start <= day <= end


def mask_name(name):
    """"RAJESH KUMAR SHARMA" -> "R*** K*** Sharma"."""
    tokens = [token for token in re.split(r"[\s.]+", str(name or "")) if token]
    if not tokens:
        return ""
    head = [token[0].upper() + "***" for token in tokens[:-1]]
    return " ".join(head + [tokens[-1][:1].upper() + tokens[-1][1:].lower()])


def mask_policy(number):
    tail = re.sub(r"[^A-Za-z0-9]", "", str(number or ""))[-4:]
    return f"****{tail}" if tail else ""


def vehicle_label(record):
    parts = [record.get("colour"), record.get("make"), record.get("model"), record.get("variant")]
    text = " ".join(str(part) for part in parts if part)
    return text[:1].upper() + text[1:]


def mobile_matches(record, wa_id):
    def digits(value):
        return re.sub(r"\D", "", str(value or ""))[-10:]
    registered = digits((record or {}).get("owner_mobile"))
    return bool(registered) and registered == digits(wa_id)


def display_date(value):
    day = _day(value)
    return day.strftime("%d %b %Y") if day else str(value or "")
